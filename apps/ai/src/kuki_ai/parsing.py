"""Turns uploaded files into citable segments.

PDFs become one segment per page. Spreadsheets become one segment per block of
rows, rendered as `REF=value` cells so the model can cite exact ranges, with the
raw cell values kept for verification and for exact sums over ranges.
"""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime

import pymupdf
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from kuki_ai.config import settings
from kuki_ai.schemas import Locator, Segment

ROWS_PER_SEGMENT = 60


class UnsupportedDocumentError(Exception):
    """The file can't be parsed. Retrying won't help."""


@dataclass
class ParsedDocument:
    filename: str
    kind: str  # "pdf" | "xlsx" | "csv"
    page_count: int | None
    segments: list[Segment]
    # Page number -> PNG, for scanned pages (and every page when images are requested).
    page_images: dict[int, bytes] = field(default_factory=dict)


def kind_for(filename: str, content_type: str | None = None) -> str:
    name = filename.lower()
    if name.endswith(".pdf") or content_type == "application/pdf":
        return "pdf"
    if name.endswith(".xlsx") or (content_type or "").endswith("spreadsheetml.sheet"):
        return "xlsx"
    if name.endswith(".csv") or content_type == "text/csv":
        return "csv"
    raise UnsupportedDocumentError(f"Unsupported file type: {filename}")


def parse_document(
    data: bytes,
    filename: str,
    content_type: str | None = None,
    *,
    all_page_images: bool = False,
) -> ParsedDocument:
    kind = kind_for(filename, content_type)
    try:
        if kind == "pdf":
            return _parse_pdf(data, filename, all_page_images)
        if kind == "xlsx":
            return _parse_xlsx(data, filename)
        return _parse_csv(data, filename)
    except UnsupportedDocumentError:
        raise
    except Exception as err:  # corrupt or encrypted files
        raise UnsupportedDocumentError(f"Could not read {filename}: {err}") from err


def _parse_pdf(data: bytes, filename: str, all_page_images: bool) -> ParsedDocument:
    doc = pymupdf.open(stream=data, filetype="pdf")
    if doc.needs_pass:
        raise UnsupportedDocumentError(f"{filename} is password protected")
    segments: list[Segment] = []
    images: dict[int, bytes] = {}
    for index, page in enumerate(doc):
        number = index + 1
        # sort=True follows the visual reading order, which keeps table rows together.
        text = page.get_text("text", sort=True).strip()
        scanned = len("".join(text.split())) < settings.min_page_chars
        if scanned or all_page_images:
            images[number] = page.get_pixmap(dpi=settings.image_dpi).tobytes("png")
        segments.append(
            Segment(ordinal=index, locator=Locator(page=number), text=text, scanned=scanned)
        )
    return ParsedDocument(filename, "pdf", doc.page_count, segments, images)


def _cell_value(value):
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        return float(value)
    return str(value).strip()


def _render(value) -> str:
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return value


def _sheet_segments(sheet_name: str, rows, start_ordinal: int) -> list[Segment]:
    """`rows` yields (row_number, [(ref, value, formula | None), ...])."""
    segments: list[Segment] = []
    lines: list[str] = []
    cells: dict[str, float | str] = {}
    first_row = last_row = None
    max_col = "A"

    def flush():
        nonlocal lines, cells, first_row
        if not lines:
            return
        cell_range = f"A{first_row}:{max_col}{last_row}"
        header = f"Sheet '{sheet_name}', rows {first_row}-{last_row}"
        segments.append(
            Segment(
                ordinal=start_ordinal + len(segments),
                locator=Locator(sheet=sheet_name, cell_range=cell_range),
                text=header + "\n" + "\n".join(lines),
                cells=cells,
            )
        )
        lines, cells, first_row = [], {}, None

    for row_number, row in rows:
        if not row:
            continue
        if first_row is None:
            first_row = row_number
        last_row = row_number
        parts = []
        for ref, value, formula in row:
            cells[ref] = value
            col = "".join(ch for ch in ref if ch.isalpha())
            if len(col) > len(max_col) or (len(col) == len(max_col) and col > max_col):
                max_col = col
            part = f"{ref}={_render(value)}"
            if formula:
                part += f" (formula {formula})"
            parts.append(part)
        lines.append(" | ".join(parts))
        if len(lines) >= ROWS_PER_SEGMENT:
            flush()
    flush()
    return segments


def _parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    # Two passes: cached values for what the reader sees, formulas to show how
    # totals were produced (a hard-coded total and a SUM can disagree).
    values_wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    formulas_wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    segments: list[Segment] = []
    for sheet_name in values_wb.sheetnames:
        values_ws = values_wb[sheet_name]
        formula_rows = formulas_wb[sheet_name].iter_rows()

        def rows(values_ws=values_ws, formula_rows=formula_rows):
            for r_index, value_row in enumerate(values_ws.iter_rows(), start=1):
                formula_row = next(formula_rows, ())
                out = []
                for c_index, cell in enumerate(value_row, start=1):
                    raw = formula_row[c_index - 1].value if c_index <= len(formula_row) else None
                    formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                    value = cell.value
                    if value is None and formula is None:
                        continue
                    ref = f"{get_column_letter(c_index)}{r_index}"
                    out.append((ref, _cell_value(value) if value is not None else "", formula))
                yield r_index, out

        segments.extend(_sheet_segments(sheet_name, rows(), len(segments)))
    if not segments:
        raise UnsupportedDocumentError(f"{filename} has no cell values")
    return ParsedDocument(filename, "xlsx", None, segments)


def _parse_csv(data: bytes, filename: str) -> ParsedDocument:
    text = data.decode("utf-8-sig", errors="replace")

    def rows():
        for r_index, row in enumerate(csv.reader(io.StringIO(text)), start=1):
            out = []
            for c_index, raw in enumerate(row, start=1):
                if raw.strip() == "":
                    continue
                try:
                    value: float | str = float(raw.replace(",", ""))
                except ValueError:
                    value = raw.strip()
                out.append((f"{get_column_letter(c_index)}{r_index}", value, None))
            yield r_index, out

    sheet = filename.rsplit(".", 1)[0]
    segments = _sheet_segments(sheet, rows(), 0)
    if not segments:
        raise UnsupportedDocumentError(f"{filename} is empty")
    return ParsedDocument(filename, "csv", None, segments)
